#!/usr/bin/env python3
"""
Anna's Archive 书目自动搜索脚本
功能：读取 books_input.csv，逐本在 Anna's Archive 搜索，输出带下载链接的 Excel

安装依赖：
    pip install playwright openpyxl
    playwright install chromium

运行方式：
    python search_anna.py
"""

import asyncio
import csv
import random
import re
from pathlib import Path
from urllib.parse import quote

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from playwright.async_api import async_playwright, TimeoutError as PlaywrightTimeout

# ── 配置 ──────────────────────────────────────────────────────────────────────
BASE_URL     = "https://annas-archive.gl"
INPUT_FILE   = "books_input.csv"
OUTPUT_FILE  = "书目搜索结果.xlsx"
PROGRESS_CSV = "search_progress.csv"   # 断点续传文件
HEADLESS     = False                    # False = 显示浏览器窗口（推荐，避免被检测）
DELAY_MIN    = 2.0                      # 每次搜索最小等待秒数
DELAY_MAX    = 4.5                      # 每次搜索最大等待秒数
SAVE_EVERY   = 50                       # 每搜索多少本保存一次进度

# ── 代理设置（科学上网）────────────────────────────────────────────────────────
# 如果你用 Clash / V2Ray / SSR 等代理工具，填写本地代理地址
# 常见默认端口：Clash = 7890，V2Ray/SSR = 1080，其他工具请在软件里查看
#
# 如果是 TUN模式 / 全局VPN（所有流量都走VPN），把下面改成 PROXY = None
#
PROXY = "http://127.0.0.1:10808"  # V2RayN 默认端口
# PROXY = None                      # 如果启用了 TUN 模式，改用这行
# ─────────────────────────────────────────────────────────────────────────────


def clean_query(title: str) -> str:
    """清理书名，提取核心关键词用于搜索"""
    # 去掉分册页码后缀，如 _1-200、_201-400
    title = re.sub(r'[\s_]+\d{1,4}[-－]\d{1,4}$', '', title)
    # 去掉 分册1、上册、下册
    title = re.sub(r'\s*(上|下|中|\d+)册$', '', title)
    title = re.sub(r'\s*第\d+[册卷分]$', '', title)
    # 去掉末尾括号里的版次 (第8版)
    title = re.sub(r'\s*[（(][第原书]\d+版[）)]$', '', title)
    # 去掉尾部数字编号 _1  _2
    title = re.sub(r'[_]\d+$', '', title)
    # 去掉作者信息（括号内含中文名或英文名）
    title = re.sub(r'\s*[（(][^）)]{2,20}著[）)]', '', title)
    title = title.strip(' _\t')

    # 超长书名只取核心部分（冒号/破折号前）
    if len(title) > 45:
        for sep in ['：', '——', ' - ', '—', '+']:
            idx = title.find(sep)
            if 4 < idx < 40:
                title = title[:idx]
                break

    return title.strip()


async def search_book(page, title: str):
    """
    在 Anna's Archive 搜索一本书。
    返回 (found_title, download_url)，找不到返回 (None, None)
    """
    query = clean_query(title)

    for lang_param in ["&lang=zh&sort=", ""]:          # 先试中文，再试全语言
        try:
            url = f"{BASE_URL}/search?q={quote(query)}{lang_param}"
            await page.goto(url, wait_until="domcontentloaded", timeout=30000)
            await asyncio.sleep(random.uniform(0.6, 1.4))

            result = await _first_result(page)
            if result:
                found_title, href = result
                detail_url = BASE_URL + href if href.startswith("/") else href

                # 进入详情页取下载链接
                await page.goto(detail_url, wait_until="domcontentloaded", timeout=30000)
                await asyncio.sleep(random.uniform(0.6, 1.2))
                dl = await _download_link(page, detail_url)
                return found_title, dl

        except PlaywrightTimeout:
            continue
        except Exception as e:
            print(f"    ⚠ {e}")
            continue

    return None, None


async def _first_result(page):
    """从搜索结果页提取第一条记录 (title, href)"""
    # 尝试多种选择器，适配网站改版
    for selector in [
        "a.js-vim-focus",
        "div.mb-4 a[href^='/md5/']",
        "a[href^='/md5/']",
    ]:
        items = await page.query_selector_all(selector)
        for item in items:
            href = await item.get_attribute("href") or ""
            if "/md5/" not in href:
                continue
            # 取书名文本
            for title_sel in ["h3", ".font-bold", "div.truncate", "div"]:
                el = await item.query_selector(title_sel)
                if el:
                    txt = (await el.inner_text()).strip().split("\n")[0]
                    if txt:
                        return txt, href
            # 没找到子元素就用 item 自身
            txt = (await item.inner_text()).strip().split("\n")[0]
            return txt, href
    return None


async def _download_link(page, fallback: str) -> str:
    """从详情页提取最优下载链接，返回 detail 页 URL 作为兜底"""
    # 先滚动确保页面内容加载
    await page.evaluate("window.scrollTo(0, document.body.scrollHeight)")
    await asyncio.sleep(0.8)

    # 明确的镜像 / 慢速下载链接（排除账户页等干扰）
    patterns = [
        "a[href*='/slow_download/']",
        "a[href*='libgen.rs']",
        "a[href*='libgen.is']",
        "a[href*='library.lol']",
        "a[href*='b-ok.']",
        "a[href*='z-lib.']",
        "a[href*='books.ms']",
        "a[href*='ipfs']",
    ]
    for pat in patterns:
        els = await page.query_selector_all(pat)
        for el in els:
            href = await el.get_attribute("href") or ""
            # 跳过账户/导航类链接
            if any(skip in href for skip in ["/account/", "/search", "#", "javascript:"]):
                continue
            if href:
                return BASE_URL + href if href.startswith("/") else href

    # 兜底：返回详情页本身（/md5/ URL），供下载脚本处理
    return fallback


# ── CSV / Excel 辅助 ──────────────────────────────────────────────────────────

FIELDS = ["序号", "分类", "原书名", "状态", "找到书名", "下载链接"]


def load_input() -> list[dict]:
    with open(INPUT_FILE, encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


def load_progress() -> dict:
    done = {}
    if Path(PROGRESS_CSV).exists():
        with open(PROGRESS_CSV, encoding="utf-8-sig") as f:
            for row in csv.DictReader(f):
                done[row["序号"]] = row
    return done


def append_progress(row: dict):
    exists = Path(PROGRESS_CSV).exists()
    with open(PROGRESS_CSV, "a", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=FIELDS)
        if not exists:
            w.writeheader()
        w.writerow({k: row.get(k, "") for k in FIELDS})


def save_excel(rows: list[dict], path: str):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "书目搜索结果"

    hdr_fill  = PatternFill("solid", fgColor="1F4E79")
    hdr_font  = Font(bold=True, color="FFFFFF", size=11)
    fill_old  = PatternFill("solid", fgColor="DDEBF7")   # 蓝：原来已找到
    fill_new  = PatternFill("solid", fgColor="E2EFDA")   # 绿：新找到
    fill_miss = PatternFill("solid", fgColor="FCE4D6")   # 红：未找到
    link_font = Font(color="0563C1", underline="single")

    headers = ["序号", "分类", "原书名", "状态", "找到书名", "下载链接"]
    widths  = [7,      18,     50,      12,     45,       14]

    for c, (h, w) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(1, c, h)
        cell.fill, cell.font = hdr_fill, hdr_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.row_dimensions[1].height = 28

    for r, book in enumerate(rows, 2):
        st = book.get("状态", "")
        dl = book.get("下载链接", "")

        if "OK" in st:
            fill = fill_old
        elif dl.startswith("http"):
            fill = fill_new
        else:
            fill = fill_miss

        vals = [book.get(k, "") for k in headers]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(r, c, v)
            cell.fill = fill
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        # 下载链接列变超链接
        if dl.startswith("http"):
            lc = ws.cell(r, 6)
            lc.hyperlink = dl
            lc.font = link_font
            lc.value = "点击下载"

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:F{len(rows)+1}"
    wb.save(path)
    print(f"\n✅ 已保存: {path}")


def build_result_rows(books, progress) -> list[dict]:
    rows = []
    for b in books:
        num = b.get("序号", "")
        if b.get("状态", "").strip() == "OK 找到":
            row = dict(b)
            row["下载链接"] = ""
            rows.append(row)
        elif num in progress:
            rows.append(progress[num])
        else:
            row = dict(b)
            row.setdefault("找到书名", "")
            row.setdefault("下载链接", "")
            rows.append(row)
    return rows


# ── 主流程 ────────────────────────────────────────────────────────────────────

async def main():
    if not Path(INPUT_FILE).exists():
        print(f"❌ 找不到 {INPUT_FILE}，请先运行 convert_to_csv.py 生成输入文件")
        return

    books    = load_input()
    progress = load_progress()

    def already_found(status: str) -> bool:
        s = status.strip()
        return "OK" in s or (s.startswith("✅") and "找到" in s)

    need_search = [b for b in books if not already_found(b.get("状态", ""))]
    pending     = [b for b in need_search if b.get("序号", "") not in progress]

    print(f"📚 总计 {len(books)} 条 | 待搜索 {len(need_search)} 条 | 本次处理 {len(pending)} 条")
    if not pending:
        print("✅ 全部已处理，直接生成 Excel …")
        save_excel(build_result_rows(books, progress), OUTPUT_FILE)
        return

    async with async_playwright() as pw:
        launch_args = ["--disable-blink-features=AutomationControlled", "--no-sandbox"]
        if PROXY:
            launch_args.append(f"--proxy-server={PROXY}")

        browser = await pw.chromium.launch(
            headless=HEADLESS,
            args=launch_args,
        )

        ctx_kwargs = dict(
            user_agent=(
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                "AppleWebKit/537.36 (KHTML, like Gecko) "
                "Chrome/121.0.0.0 Safari/537.36"
            ),
            viewport={"width": 1366, "height": 800},
            locale="zh-CN",
        )
        if PROXY:
            ctx_kwargs["proxy"] = {"server": PROXY}

        ctx = await browser.new_context(**ctx_kwargs)
        page = await ctx.new_page()

        if PROXY:
            print(f"🌐 代理已启用: {PROXY}")
        else:
            print("🌐 未配置代理（TUN全局模式 或 直连）")

        print(f"\n🔗 打开 {BASE_URL} …")
        try:
            await page.goto(BASE_URL, wait_until="domcontentloaded", timeout=40000)
            await asyncio.sleep(3)
            print("✅ 网站已加载，开始搜索\n")
        except Exception as e:
            print(f"⚠ 首页加载异常: {e}，继续尝试搜索…")

        found_count = 0
        for i, book in enumerate(pending):
            num   = book.get("序号", "?")
            title = book.get("原书名", "")
            print(f"[{i+1:4d}/{len(pending)}] #{num:>4}  {title[:50]}")

            found_title, dl_link = await search_book(page, title)

            row = dict(book)
            if found_title:
                row["状态"]     = "✅ 已找到"
                row["找到书名"] = found_title
                row["下载链接"] = dl_link or ""
                found_count += 1
                print(f"          ✅ {found_title[:55]}")
            else:
                row["状态"]     = "❌ 未找到"
                row["找到书名"] = ""
                row["下载链接"] = ""
                print(f"          ❌ 未找到")

            progress[num] = row
            append_progress(row)

            if (i + 1) % SAVE_EVERY == 0:
                mid_path = f"进度_{i+1}.xlsx"
                save_excel(build_result_rows(books, progress), mid_path)

            await asyncio.sleep(random.uniform(DELAY_MIN, DELAY_MAX))

        await browser.close()

    all_rows = build_result_rows(books, progress)
    save_excel(all_rows, OUTPUT_FILE)
    print(f"\n🎉 完成！新找到 {found_count}/{len(pending)} 本 | Excel: {OUTPUT_FILE}")


if __name__ == "__main__":
    asyncio.run(main())
