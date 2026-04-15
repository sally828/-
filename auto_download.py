#!/usr/bin/env python3
"""
Anna's Archive 自动下载脚本（Playwright 版）
- 浏览器打开后你手动登录一次，后续全自动
- 运行：python auto_download.py
"""

import asyncio
import random
import re
from pathlib import Path

import openpyxl
import requests as _req
from playwright.async_api import async_playwright, TimeoutError as PlaywrightTimeout

# ── 配置 ──────────────────────────────────────────────
PROXY        = "http://127.0.0.1:10808"   # TUN全局模式改成 None
HEADLESS     = False
DELAY_MIN    = 1.5
DELAY_MAX    = 3.0
BASE_URL     = "https://annas-archive.gl"
# ─────────────────────────────────────────────────────

_HERE        = Path(__file__).parent
EXCEL_FILE   = _HERE / "书目搜索结果.xlsx"
DOWNLOAD_DIR = _HERE / "下载书籍"
DONE_FILE    = _HERE / "downloaded.txt"

# 需要拦截的书籍文件类型
FILE_TYPES = ["pdf", "epub", "mobi", "djvu", "azw3", "azw", "fb2", "cbz", "cbr"]


def extract_md5(url: str) -> str | None:
    m = re.search(r'[a-f0-9]{32}', str(url or ""), re.I)
    return m.group(0).lower() if m else None

def safe_name(name: str) -> str:
    return re.sub(r'[\\/:*?"<>|]', '_', name)[:80].strip()

def load_done() -> set:
    return set(DONE_FILE.read_text(encoding="utf-8").splitlines()) if DONE_FILE.exists() else set()

def mark_done(num: str):
    with open(DONE_FILE, "a", encoding="utf-8") as f:
        f.write(num + "\n")

def _do_download(cdn_url: str, save_dir: Path, base_name: str, proxies) -> Path:
    """同步下载函数，在线程中运行，字节直接写盘，不经过 Playwright IPC"""
    headers = {
        "User-Agent": (
            "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
            "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/121.0.0.0 Safari/537.36"
        )
    }
    r = _req.get(cdn_url, headers=headers, proxies=proxies, timeout=120, stream=True)
    r.raise_for_status()
    ct  = r.headers.get("content-type", "").lower()
    ext = next((t for t in FILE_TYPES if t in ct or cdn_url.lower().endswith(f".{t}")), "pdf")
    sp   = save_dir / f"{base_name}.{ext}"
    size = 0
    try:
        with open(sp, "wb") as f:
            for chunk in r.iter_content(chunk_size=65536):
                if chunk:
                    f.write(chunk)
                    size += len(chunk)
    except Exception:
        sp.unlink(missing_ok=True)   # 下载中断时清理残缺文件
        raise
    if size < 10_000:
        sp.unlink(missing_ok=True)
        raise ValueError(f"文件太小（{size} 字节）")
    return sp


def load_books() -> list[dict]:
    wb = openpyxl.load_workbook(EXCEL_FILE)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=False):
        data = {}
        for i, cell in enumerate(row):
            if i < len(headers) and headers[i]:
                data[headers[i]] = cell.hyperlink.target if cell.hyperlink else cell.value
        num  = str(data.get("序号", "") or "").strip()
        name = str(data.get("找到书名", "") or data.get("原书名", "") or "").strip()
        url  = str(data.get("下载链接", "") or "").strip()
        if num:
            rows.append({"序号": num, "书名": name or f"书_{num}", "url": url})
    return rows


async def download_book(page, book: dict) -> bool:
    url   = book["url"]
    title = book["书名"]

    if not url.startswith("http"):
        print(f"    ⚠ 无链接，跳过")
        return False

    # ── 第一步：导航到 fast_download URL，让浏览器跟随跳转到 CDN ──
    try:
        await page.goto(url, wait_until="commit", timeout=30000)
    except Exception:
        pass

    # ── 第二步：等待浏览器地址栏变成 CDN 地址（最多 12 秒）──
    cdn_url = url
    for _ in range(12):
        await asyncio.sleep(1)
        current = page.url
        # 排除 chrome-error:// 等非 http 地址（CDN 无法访问时浏览器跳到内部错误页）
        if (current
                and current.startswith("http")
                and "annas-archive" not in current
                and current != url):
            cdn_url = current
            break

    if not cdn_url.startswith("http") or "annas-archive" in cdn_url or cdn_url == url:
        print(f"    ❌ 未能跳转到下载地址（CDN 不可达或代理断线）")
        return False

    # ── 第三步：requests 在线程里直接下载到磁盘 ──
    # 不走 Playwright IPC——大文件（100MB+）会撑爆 IPC Socket
    proxies = {"http": PROXY, "https": PROXY} if PROXY else None
    try:
        sp = await asyncio.to_thread(
            _do_download, cdn_url, DOWNLOAD_DIR, safe_name(title), proxies
        )
        size_mb = sp.stat().st_size / 1024 / 1024
        print(f"    ✅ {sp.name}  ({size_mb:.1f} MB)")
        return True
    except Exception as e:
        print(f"    ❌ {e}")
        return False


async def main():
    DOWNLOAD_DIR.mkdir(exist_ok=True)
    books   = load_books()
    done    = load_done()
    pending = [b for b in books if b["序号"] not in done and extract_md5(b["url"])]

    print(f"📚 共 {len(books)} 条 | 已完成 {len(done)} | 本次下载 {len(pending)} 条")
    print(f"📁 保存到: {DOWNLOAD_DIR}\n")
    if not pending:
        print("✅ 全部完成！")
        return

    async with async_playwright() as pw:
        args = [
            "--disable-blink-features=AutomationControlled",
            "--no-sandbox",
        ]
        if PROXY:
            args.append(f"--proxy-server={PROXY}")

        browser = await pw.chromium.launch(headless=HEADLESS, args=args)
        ctx = await browser.new_context(
            user_agent=(
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                "AppleWebKit/537.36 (KHTML, like Gecko) "
                "Chrome/121.0.0.0 Safari/537.36"
            ),
            viewport={"width": 1366, "height": 800},
            locale="zh-CN",
            accept_downloads=True,
        )
        page = await ctx.new_page()

        print("🌐 正在打开 Anna's Archive…")
        await page.goto(f"{BASE_URL}/account", wait_until="domcontentloaded", timeout=40000)
        print("\n" + "=" * 55)
        print("请在浏览器里登录你的 Anna's Archive 账号")
        print("登录完成后回到这里按 Enter")
        print("=" * 55)
        input(">>> 登录完成，按 Enter 开始下载：")
        print()

        success = fail = 0
        for i, book in enumerate(pending):
            print(f"[{i+1:4d}/{len(pending)}] #{book['序号']}  {book['书名'][:50]}")
            ok = await download_book(page, book)
            if ok:
                success += 1
                mark_done(book["序号"])
            else:
                fail += 1
            await asyncio.sleep(random.uniform(DELAY_MIN, DELAY_MAX))

        await browser.close()

    print(f"\n🎉 完成！成功 {success} | 失败 {fail}")


if __name__ == "__main__":
    asyncio.run(main())
