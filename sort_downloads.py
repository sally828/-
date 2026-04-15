#!/usr/bin/env python3
"""
下载文件分拣脚本
读取 verify.py 生成的「核验报告.xlsx」，自动：
  1. 把文件匹配度 < 60% 或备注含「差异较大」的文件 → 移到「待删除」文件夹
  2. 从 downloaded.txt 删除被移走文件的序号（让下次重跑时重新下载）
  3. 输出「待下载清单.txt」（⏳未下载 + 被移走的书目）

运行：python sort_downloads.py
前提：先运行 verify.py 生成核验报告.xlsx
"""

import shutil
from pathlib import Path
from tkinter import Tk, filedialog, messagebox

import openpyxl

_HERE     = Path(__file__).parent
DONE_FILE = _HERE / "downloaded.txt"
TODO_FILE = _HERE / "待下载清单.txt"


# ── 工具函数 ──────────────────────────────────────────────────────────────────

def pick_file(title: str, filetypes=(("Excel", "*.xlsx"), ("所有文件", "*.*"))) -> str | None:
    root = Tk(); root.withdraw(); root.attributes("-topmost", True)
    path = filedialog.askopenfilename(title=title, filetypes=filetypes)
    root.destroy()
    return path or None


def pick_folder(title: str) -> Path | None:
    root = Tk(); root.withdraw(); root.attributes("-topmost", True)
    path = filedialog.askdirectory(title=title)
    root.destroy()
    return Path(path) if path else None


def is_bad_match(note: str, file_sim) -> bool:
    """判断是否为错配文件"""
    if "差异较大" in str(note):
        return True
    try:
        return float(file_sim) < 60
    except (TypeError, ValueError):
        return False


def move_file(src: Path, trash_folder: Path) -> Path:
    """移动文件到待删除文件夹，自动处理同名冲突"""
    dst = trash_folder / src.name
    counter = 1
    while dst.exists():
        dst = trash_folder / f"{src.stem}_{counter}{src.suffix}"
        counter += 1
    shutil.move(str(src), str(dst))
    return dst


# ── 主流程 ────────────────────────────────────────────────────────────────────

def main():
    root = Tk(); root.withdraw()
    messagebox.showinfo(
        "下载文件分拣脚本",
        "接下来依次选择：\n\n"
        "① 核验报告.xlsx（由 verify.py 生成）\n"
        "② 下载书籍所在文件夹\n\n"
        "脚本会把错配文件移到「待删除」文件夹，\n"
        "并生成「待下载清单.txt」。"
    )
    root.destroy()

    report_path = pick_file("① 选择核验报告.xlsx")
    if not report_path:
        print("已取消"); return

    dl_folder = pick_folder("② 选择下载书籍所在文件夹")
    if not dl_folder:
        print("已取消"); return

    trash_folder = dl_folder.parent / "待删除"
    trash_folder.mkdir(exist_ok=True)

    # ── 读取核验报告 ──
    wb = openpyxl.load_workbook(report_path)
    ws = wb.active
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]

    def idx(name: str) -> int | None:
        return headers.index(name) if name in headers else None

    i_num   = idx("序号")
    i_orig  = idx("原书名")
    i_file  = idx("下载文件名")
    i_fsim  = idx("文件匹配度%")
    i_note  = idx("备注")

    if any(v is None for v in [i_num, i_orig, i_file, i_fsim, i_note]):
        print("❌ 核验报告列名不匹配，请确认是 verify.py 生成的文件")
        return

    # ── 读取 downloaded.txt ──
    done_nums: set[str] = set()
    if DONE_FILE.exists():
        done_nums = set(DONE_FILE.read_text(encoding="utf-8").splitlines())

    moved_nums:   set[str]  = set()
    need_download: list[str] = []

    moved = skipped_notfound = already_ok = 0

    print(f"📊 读取核验报告：{report_path}\n")

    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[i_num] is None:
            continue

        num   = str(row[i_num] or "").strip()
        orig  = str(row[i_orig] or "").strip()
        fname = str(row[i_file] or "").strip()
        fsim  = row[i_fsim]
        note  = str(row[i_note] or "").strip()

        # ⏳ 未下载 → 直接加入待下载清单
        if "⏳" in note:
            need_download.append(f"{num}\t{orig}")
            continue

        # ❌ Anna 未收录 → 跳过（无法下载）
        if "❌" in note:
            skipped_notfound += 1
            continue

        # ✅ 正常 → 不动
        if "✅" in note and not is_bad_match(note, fsim):
            already_ok += 1
            continue

        # ⚠ 有文件但错配 → 移到待删除
        if fname and is_bad_match(note, fsim):
            src = dl_folder / fname
            if src.exists():
                dst = move_file(src, trash_folder)
                print(f"  📦 移走：{fname[:55]}")
                print(f"       → {dst.name[:55]}")
                moved += 1
            else:
                print(f"  ⚠ 文件不存在（可能已删除）：{fname[:55]}")
                moved += 1          # 也算进去，从 downloaded.txt 里删掉

            moved_nums.add(num)
            need_download.append(f"{num}\t{orig}")

    # ── 更新 downloaded.txt ──
    if moved_nums and DONE_FILE.exists():
        remaining = [n for n in done_nums if n not in moved_nums]
        DONE_FILE.write_text("\n".join(remaining) + ("\n" if remaining else ""),
                             encoding="utf-8")
        print(f"\n  🗑  已从 downloaded.txt 删除 {len(moved_nums)} 个序号")

    # ── 输出待下载清单 ──
    with open(TODO_FILE, "w", encoding="utf-8") as f:
        f.write("# 待下载书目（⏳未下载 + ⚠错配被移走）\n")
        f.write(f"# 共 {len(need_download)} 条\n")
        f.write("# 格式：序号<Tab>原书名\n\n")
        for line in need_download:
            f.write(line + "\n")

    summary = (
        f"📦 移到待删除：{moved} 个\n"
        f"✅ 正常保留：{already_ok} 个\n"
        f"❌ Anna 未收录（不可下载）：{skipped_notfound} 个\n"
        f"📋 待下载清单：{len(need_download)} 条\n\n"
        f"待删除文件夹：{trash_folder}\n"
        f"待下载清单：{TODO_FILE}\n\n"
        f"确认「待删除」内容无误后可手动清空该文件夹，\n"
        f"然后重跑 auto_download.py 补下缺失书目。"
    )
    print(f"\n🎉 分拣完成！\n{summary}")

    root = Tk(); root.withdraw()
    messagebox.showinfo("分拣完成", summary)
    root.destroy()


if __name__ == "__main__":
    main()
